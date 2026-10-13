"""Document parsing into structural blocks that keep page numbers and section headings.

PDF: Docling (layout, reading order, tables as Markdown) when installed, otherwise pypdf.
Also handles .docx, .md and .txt.
"""

from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path

from ragdoc.types import Block, ParsedDocument

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _doc_id(path: Path, sha: str) -> str:
    stem = re.sub(r"[^a-zA-Z0-9]+", "-", path.stem).strip("-").lower()[:40] or "doc"
    return f"{stem}-{sha[:8]}"


def _clean(text: str) -> str:
    text = text.replace("­", "")  # soft hyphens
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- PDF: docling


def docling_available() -> bool:
    try:
        import docling  # noqa: F401
    except ImportError:
        return False
    return True


def parse_pdf_docling(path: Path) -> tuple[list[Block], int | None]:
    from docling.document_converter import DocumentConverter

    result = DocumentConverter().convert(str(path))
    doc = result.document
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []  # (level, text) stack

    for item, _depth in doc.iterate_items():
        label = str(getattr(item, "label", "")).lower()
        prov = getattr(item, "prov", None) or []
        page = getattr(prov[0], "page_no", None) if prov else None

        if "table" in label and hasattr(item, "export_to_markdown"):
            try:
                text = item.export_to_markdown(doc=doc)
            except TypeError:  # older docling signature
                text = item.export_to_markdown()
            kind = "table"
        else:
            text = getattr(item, "text", "") or ""
            kind = "text"
            if "section_header" in label or label.endswith("title"):
                level = int(getattr(item, "level", 1) or 1)
                headings = [h for h in headings if h[0] < level] + [(level, text.strip())]
                kind = "heading"
            elif "list" in label:
                kind = "list"
            elif "caption" in label:
                kind = "caption"
            elif label.endswith(("page_header", "page_footer")):
                continue  # running headers/footers add noise

        text = _clean(text)
        if not text:
            continue
        blocks.append(Block(text=text, page=page, kind=kind, heading_path=[h[1] for h in headings]))

    num_pages = len(getattr(doc, "pages", {}) or {}) or None
    return blocks, num_pages


# --------------------------------------------------------------------------- PDF: pypdf


_HEADING_RE = re.compile(r"^(\d+(\.\d+)*\.?\s+)?[A-Z][^.!?]{2,80}$")


def parse_pdf_pypdf(path: Path) -> tuple[list[Block], int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    heading_path: list[str] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        for para in re.split(r"\n\s*\n", text):
            para = para.strip()
            if not para:
                continue
            # Heuristic: short single line, title-like → heading.
            if "\n" not in para and len(para) < 90 and _HEADING_RE.match(para):
                heading_path = [para]
                blocks.append(Block(text=para, page=page_no, kind="heading", heading_path=[para]))
                continue
            blocks.append(
                Block(text=para.replace("\n", " "), page=page_no, heading_path=list(heading_path))
            )
    return blocks, len(reader.pages)


# --------------------------------------------------------------------------- other formats


def parse_docx(path: Path) -> list[Block]:
    import docx

    document = docx.Document(str(path))
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []
    for para in document.paragraphs:
        text = _clean(para.text)
        if not text:
            continue
        style = (para.style.name if para.style is not None else "").lower()
        if style.startswith("heading") or style == "title":
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
            headings = [h for h in headings if h[0] < level] + [(level, text)]
            blocks.append(Block(text=text, kind="heading", heading_path=[h[1] for h in headings]))
        else:
            kind = "list" if "list" in style else "text"
            blocks.append(Block(text=text, kind=kind, heading_path=[h[1] for h in headings]))
    for table in document.tables:
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        header, *body = rows
        md = "| " + " | ".join(header) + " |\n| " + " | ".join("---" for _ in header) + " |\n"
        md += "\n".join("| " + " | ".join(r) + " |" for r in body)
        blocks.append(Block(text=md, kind="table"))
    return blocks


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []
    buffer: list[str] = []

    def flush() -> None:
        para = _clean("\n".join(buffer))
        buffer.clear()
        if para:
            kind = "table" if para.lstrip().startswith("|") else "text"
            blocks.append(Block(text=para, kind=kind, heading_path=[h[1] for h in headings]))

    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if match:
            flush()
            level, title = len(match.group(1)), match.group(2).strip()
            headings = [h for h in headings if h[0] < level] + [(level, title)]
            blocks.append(Block(text=title, kind="heading", heading_path=[h[1] for h in headings]))
        elif not line.strip():
            flush()
        else:
            buffer.append(line)
    flush()
    return blocks


# --------------------------------------------------------------------------- entry point


def parse_document(path: str | Path, parser: str = "auto") -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Unsupported file type {suffix!r}; supported: {sorted(SUPPORTED_SUFFIXES)}"
        )
    sha = file_sha256(path)
    num_pages: int | None = None

    if suffix == ".pdf":
        use_docling = parser == "docling" or (parser == "auto" and docling_available())
        if use_docling:
            try:
                blocks, num_pages = parse_pdf_docling(path)
            except Exception as exc:  # docling failures should not block ingestion
                if parser == "docling":
                    raise
                log.warning("Docling failed on %s (%s); falling back to pypdf", path.name, exc)
                blocks, num_pages = parse_pdf_pypdf(path)
        else:
            blocks, num_pages = parse_pdf_pypdf(path)
    elif suffix == ".docx":
        blocks = parse_docx(path)
    else:
        blocks = parse_markdown(path.read_text(encoding="utf-8", errors="replace"))

    if not blocks:
        raise ValueError(
            f"No text could be extracted from {path.name} (scanned PDF? try OCR first)"
        )

    title = next((b.text for b in blocks if b.kind == "heading"), path.stem)
    return ParsedDocument(
        doc_id=_doc_id(path, sha),
        source=path.name,
        title=title[:200],
        blocks=blocks,
        num_pages=num_pages,
        sha256=sha,
    )
